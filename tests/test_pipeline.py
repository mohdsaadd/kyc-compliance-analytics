from datetime import date

import pandas as pd
import pytest
from openpyxl import load_workbook

from main import run_pipeline
from src.clean import clean_cases
from src.database import load_cases, load_queries, run_query

AS_OF = date(2026, 10, 27)


def raw_df(**overrides):
    base = {"case_id": "KYC-1", "customer_name": "  aarav sharma", "document_type": "Passport",
            "region": "north ", "risk_level": "high", "status": " Done ", "analyst": "Analyst A",
            "submitted_date": "2026-10-01", "completed_date": "2026-10-04"}
    base.update(overrides)
    return pd.DataFrame([base])


# ---------- cleaning ----------
def test_text_and_status_are_standardised():
    clean, issues = clean_cases(raw_df())
    row = clean.iloc[0]
    assert row["customer_name"] == "Aarav Sharma" and row["region"] == "North"
    assert row["risk_level"] == "High" and row["status"] == "Approved"
    assert issues.empty


@pytest.mark.parametrize("variant,expected", [
    ("in progress", "Pending"), ("DECLINED", "Rejected"), ("approved", "Approved")])
def test_status_variants(variant, expected):
    completed = None if expected == "Pending" else "2026-10-04"
    clean, _ = clean_cases(raw_df(status=variant, completed_date=completed))
    assert clean.iloc[0]["status"] == expected


def test_duplicates_removed_keeping_first():
    df = pd.concat([raw_df(), raw_df(customer_name="Other Person")], ignore_index=True)
    clean, issues = clean_cases(df)
    assert len(clean) == 1 and clean.iloc[0]["customer_name"] == "Aarav Sharma"
    assert issues["reason"].tolist() == ["duplicate case_id"]


def test_invalid_rows_are_logged_not_silently_dropped():
    df = pd.concat([
        raw_df(case_id="A", submitted_date="N/A"),
        raw_df(case_id="B", completed_date="2026-09-01"),            # before submitted
        raw_df(case_id="C", status="mystery"),
        raw_df(case_id="D", status="Pending", completed_date="2026-10-04"),
        raw_df(case_id="E", completed_date=None),                    # closed, no completion
        raw_df(case_id="F"),                                         # valid
    ], ignore_index=True)
    clean, issues = clean_cases(df)
    assert clean["case_id"].tolist() == ["F"]
    assert dict(zip(issues["case_id"], issues["reason"])) == {
        "A": "invalid or missing submitted_date",
        "B": "completed_date earlier than submitted_date",
        "C": "unrecognised status",
        "D": "pending case has a completed_date",
        "E": "closed case has no completed_date"}


def test_missing_columns_raise():
    with pytest.raises(ValueError, match="missing columns"):
        clean_cases(pd.DataFrame({"case_id": ["x"]}))


# ---------- SQL ----------
@pytest.fixture
def conn():
    rows = [
        raw_df(case_id="1", document_type="Passport", status="Approved", completed_date="2026-10-03",
               submitted_date="2026-10-01"),
        raw_df(case_id="2", document_type="Passport", status="Approved", completed_date="2026-10-07",
               submitted_date="2026-10-01"),
        raw_df(case_id="3", document_type="PAN Card", status="Pending", completed_date=None,
               submitted_date="2026-10-20", risk_level="Low"),                         # 7 days old
        raw_df(case_id="4", document_type="PAN Card", status="Pending", completed_date=None,
               submitted_date="2026-10-10", risk_level="High"),                        # 17 days old
        raw_df(case_id="5", document_type="PAN Card", status="Pending", completed_date=None,
               submitted_date="2026-10-26"),                                           # 1 day old
    ]
    clean, _ = clean_cases(pd.concat(rows, ignore_index=True))
    return load_cases(clean)


def test_all_named_queries_load():
    assert set(load_queries()) == {"status_summary", "turnaround_by_document", "analyst_performance",
                                   "monthly_trend", "overdue_cases"}


def test_turnaround_average(conn):
    df = run_query(conn, "turnaround_by_document")
    assert df.iloc[0]["document_type"] == "Passport"
    assert df.iloc[0]["avg_turnaround_days"] == 4.0          # (2 + 6) / 2
    assert df.iloc[0]["max_turnaround_days"] == 6


def test_overdue_respects_sla_and_sorts_high_risk_first(conn):
    df = run_query(conn, "overdue_cases", as_of=AS_OF.isoformat(), sla_days=5)
    assert df["case_id"].tolist() == ["4", "3"]               # case 5 is within SLA
    assert df["days_pending"].tolist() == [17, 7]
    assert run_query(conn, "overdue_cases", as_of=AS_OF.isoformat(), sla_days=30).empty


def test_unknown_query_name(conn):
    with pytest.raises(KeyError):
        run_query(conn, "nope")


def test_status_summary_percentages_add_up(conn):
    df = run_query(conn, "status_summary")
    assert df["cases"].sum() == 5 and round(df["pct_of_total"].sum()) == 100


# ---------- end to end ----------
def test_pipeline_creates_report_with_live_formulas(tmp_path):
    report = run_pipeline(output_dir=tmp_path, as_of=AS_OF, sla_days=5, rows=120)
    wb = load_workbook(report)
    assert {"Summary", "Cases", "Overdue_Cases", "Data_Issues"} <= set(wb.sheetnames)
    summary = wb["Summary"]
    assert summary["B8"].value.startswith("=COUNTA(")          # formulas, not hardcoded results
    assert "COUNTIFS" in summary["B14"].value
    assert summary["B4"].value.date() == AS_OF and summary["B5"].value == 5
    assert (tmp_path / "kyc.sqlite").exists() and (tmp_path / "raw_kyc_export.xlsx").exists()
    # one Cases row per cleaned record
    assert wb["Cases"].max_row - 1 == len(pd.read_excel(report, sheet_name="Cases"))
